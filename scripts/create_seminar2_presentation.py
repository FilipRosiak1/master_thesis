from __future__ import annotations

import csv
import math
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
CPU_RUN = ROOT / "exports" / "final_results_cpu" / "exports" / "final_seed_bucket_benchmark_baseline_20260613_011516"
GPU_RUN = ROOT / "exports" / "final_results_gpu" / "exports" / "final_seed_bucket_benchmark_20260607_194941"
OUT_DIR = ROOT / "seminar"
ASSET_DIR = OUT_DIR / "wystapienie2_assets"
PPTX_PATH = OUT_DIR / "wystapienie2_latent_framsticks.pptx"
NOTES_PATH = OUT_DIR / "wystapienie2_tekst_do_slajdow.md"


FITNESS_AWARE = "exports/cluster_export_20260524_222538/models/f1_guided/fitness_aware_lhs_seed42"
LABEL_NAMES = {
    "09_tree_vae_epoch80": "TreeVAE epoch80",
    "10_transformer_vae_epoch90": "TransformerVAE epoch90",
    "02_lhs_latent256_seed42_best": "LHS 256",
    "06_lhs_depth_seed42_best": "LHS depth",
    FITNESS_AWARE: "Fitness-aware LHS",
    "": "Framsticks EA",
}
METHOD_NAMES = {
    "frams": "Framsticks EA",
    "latent_ea": "EA w przestrzeni ukrytej",
    "cem": "CEM",
    "cmaes": "CMA-ES",
}
BUCKET_ORDER = ["B1", "B2", "B3", "B4", "B5", "B6"]
BUCKET_LABELS = {
    "B1": "0.30-0.50",
    "B2": "0.50-0.75",
    "B3": "0.75-1.00",
    "B4": "1.00-1.20",
    "B5": "1.20-1.50",
    "B6": ">=1.50",
}


BG = RGBColor(11, 16, 32)
BG_ALT = RGBColor(17, 24, 39)
TEXT = RGBColor(245, 247, 250)
MUTED = RGBColor(173, 184, 201)
LINE = RGBColor(71, 85, 105)
CYAN = RGBColor(45, 212, 191)
BLUE = RGBColor(56, 189, 248)
AMBER = RGBColor(245, 158, 11)
VIOLET = RGBColor(167, 139, 250)
GREEN = RGBColor(52, 211, 153)
PINK = RGBColor(244, 114, 182)
RED = RGBColor(248, 113, 113)


@dataclass
class SlideText:
    title: str
    speaker: str


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def load_job_rows(run_dir: Path, subdir: str, *, history: bool = False) -> list[dict[str, str]]:
    suffix = "*.history.csv" if history else "*.csv"
    rows: list[dict[str, str]] = []
    for path in sorted((run_dir / "jobs" / subdir).glob(suffix)):
        if not history and path.name.endswith(".history.csv"):
            continue
        rows.extend(read_csv(path))
    return rows


def fl(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return float("nan")


def fmt(value: float, digits: int = 3) -> str:
    return f"{value:.{digits}f}"


def pct(value: float, digits: int = 0) -> str:
    return f"{100 * value:.{digits}f}%"


def mean(values: list[float]) -> float:
    values = [value for value in values if not math.isnan(value)]
    return sum(values) / len(values) if values else float("nan")


def median(values: list[float]) -> float:
    values = sorted(value for value in values if not math.isnan(value))
    if not values:
        return float("nan")
    mid = len(values) // 2
    return values[mid] if len(values) % 2 else 0.5 * (values[mid - 1] + values[mid])


def quantile(values: list[float], q: float) -> float:
    values = sorted(value for value in values if not math.isnan(value))
    if not values:
        return float("nan")
    idx = (len(values) - 1) * q
    lo = math.floor(idx)
    hi = math.ceil(idx)
    if lo == hi:
        return values[lo]
    return values[lo] * (hi - idx) + values[hi] * (idx - lo)


def seed_key(row: dict[str, str]) -> tuple[str, str, str]:
    return (row.get("bucket_id", ""), row.get("seed_rank", ""), row.get("seed_dataset_idx", ""))


def summary_key(row: dict[str, str]) -> tuple[str, str, str, str, str, str]:
    return (
        row.get("label", ""),
        row.get("method", ""),
        row.get("bucket_id", ""),
        row.get("seed_rank", ""),
        row.get("seed_dataset_idx", ""),
        row.get("model", ""),
    )


def label_name(label: str) -> str:
    if label in LABEL_NAMES:
        return LABEL_NAMES[label]
    return Path(label.replace("\\", "/")).name if "/" in label or "\\" in label else label


def group_stats(rows: list[dict[str, str]], keys: list[str]) -> dict[tuple[str, ...], dict[str, float]]:
    groups: dict[tuple[str, ...], list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        groups[tuple(row.get(key, "") for key in keys)].append(row)
    stats: dict[tuple[str, ...], dict[str, float]] = {}
    for key, group in groups.items():
        scores = [fl(row["best_score"]) for row in group]
        deltas = [fl(row["best_minus_seed_source"]) for row in group]
        seeds = [fl(row["seed_source_fitness"]) for row in group]
        stats[key] = {
            "n": float(len(group)),
            "max": max(scores),
            "mean": mean(scores),
            "median": median(scores),
            "mean_delta": mean(deltas),
            "median_delta": median(deltas),
            "improvement_rate": sum(score > seed + 1e-9 for score, seed in zip(scores, seeds)) / len(group),
        }
    return stats


def filter_rows(rows: list[dict[str, str]], *, label: str | None = None, method: str | None = None, bucket: str | None = None) -> list[dict[str, str]]:
    out = rows
    if label is not None:
        out = [row for row in out if row.get("label", "") == label]
    if method is not None:
        out = [row for row in out if row.get("method", "") == method]
    if bucket is not None:
        out = [row for row in out if row.get("bucket_id", "") == bucket]
    return out


def best_by_seed(rows: list[dict[str, str]], *, method: str | None = None) -> dict[tuple[str, str, str], dict[str, str]]:
    best: dict[tuple[str, str, str], dict[str, str]] = {}
    for row in rows:
        if method is not None and row.get("method") != method:
            continue
        key = seed_key(row)
        if key not in best or fl(row["best_score"]) > fl(best[key]["best_score"]):
            best[key] = row
    return best


def compare_to_baseline(rows: list[dict[str, str]], baseline: dict[tuple[str, str, str], dict[str, str]]) -> dict[tuple[str, str], dict[str, float]]:
    grouped: dict[tuple[str, str], list[float]] = defaultdict(list)
    for row in rows:
        base = baseline.get(seed_key(row))
        if base is None:
            continue
        grouped[(row.get("label", ""), row.get("method", ""))].append(fl(row["best_score"]) - fl(base["best_score"]))
    out: dict[tuple[str, str], dict[str, float]] = {}
    for key, diffs in grouped.items():
        out[key] = {
            "n": float(len(diffs)),
            "wins": float(sum(diff > 1e-9 for diff in diffs)),
            "losses": float(sum(diff < -1e-9 for diff in diffs)),
            "win_rate": sum(diff > 1e-9 for diff in diffs) / len(diffs),
            "mean_diff": mean(diffs),
            "median_diff": median(diffs),
        }
    return out


def at_step(history_rows: list[dict[str, str]], step_limit: int) -> list[dict[str, str]]:
    latest: dict[tuple[str, str, str, str, str, str], dict[str, str]] = {}
    for row in history_rows:
        step = int(fl(row.get("step") or row.get("generation") or row.get("iteration") or 0))
        if step > step_limit:
            continue
        key = summary_key(row)
        prev = latest.get(key)
        if prev is None or step > int(fl(prev.get("step") or prev.get("generation") or prev.get("iteration") or 0)):
            latest[key] = row
    return list(latest.values())


def compute_summary() -> dict[str, Any]:
    cpu_rows = load_job_rows(CPU_RUN, "cpu", history=False)
    gpu_rows = load_job_rows(GPU_RUN, "gpu", history=False)
    cpu_hist = load_job_rows(CPU_RUN, "cpu", history=True)
    gpu_hist = load_job_rows(GPU_RUN, "gpu", history=True)
    baseline = {seed_key(row): row for row in cpu_rows}
    comp = compare_to_baseline(gpu_rows, baseline)
    cpu_stats = group_stats(cpu_rows, ["method"])[("frams",)]
    gpu_method_stats = group_stats(gpu_rows, ["method"])
    gpu_label_method_stats = group_stats(gpu_rows, ["label", "method"])
    gpu_label_stats = group_stats(gpu_rows, ["label"])

    oracle_all = best_by_seed(gpu_rows)
    oracle_latent_ea = best_by_seed(gpu_rows, method="latent_ea")
    oracle_cmaes = best_by_seed(gpu_rows, method="cmaes")
    oracle_cem = best_by_seed(gpu_rows, method="cem")

    def oracle_stats(items: dict[tuple[str, str, str], dict[str, str]]) -> dict[str, float]:
        diffs: list[float] = []
        scores: list[float] = []
        for key, row in items.items():
            base = baseline[key]
            scores.append(fl(row["best_score"]))
            diffs.append(fl(row["best_score"]) - fl(base["best_score"]))
        return {
            "n": float(len(scores)),
            "wins": float(sum(diff > 1e-9 for diff in diffs)),
            "losses": float(sum(diff < -1e-9 for diff in diffs)),
            "win_rate": sum(diff > 1e-9 for diff in diffs) / len(diffs),
            "max": max(scores),
            "mean": mean(scores),
            "median": median(scores),
            "mean_diff": mean(diffs),
            "median_diff": median(diffs),
        }

    hist100_cpu = at_step(cpu_hist, 100)
    hist100_gpu = at_step(gpu_hist, 100)
    comp100 = compare_to_baseline(hist100_gpu, {seed_key(row): row for row in hist100_cpu})
    gpu_label_method_100 = group_stats(hist100_gpu, ["label", "method"])
    cpu100_stats = group_stats(hist100_cpu, ["method"])[("frams",)]

    return {
        "cpu_rows": cpu_rows,
        "gpu_rows": gpu_rows,
        "cpu_hist": cpu_hist,
        "gpu_hist": gpu_hist,
        "baseline": baseline,
        "comp": comp,
        "cpu_stats": cpu_stats,
        "gpu_method_stats": gpu_method_stats,
        "gpu_label_method_stats": gpu_label_method_stats,
        "gpu_label_stats": gpu_label_stats,
        "oracle_all": oracle_stats(oracle_all),
        "oracle_latent_ea": oracle_stats(oracle_latent_ea),
        "oracle_cmaes": oracle_stats(oracle_cmaes),
        "oracle_cem": oracle_stats(oracle_cem),
        "hist100_cpu": hist100_cpu,
        "hist100_gpu": hist100_gpu,
        "comp100": comp100,
        "gpu_label_method_100": gpu_label_method_100,
        "cpu100_stats": cpu100_stats,
    }


def setup_plot_style() -> None:
    plt.rcParams.update(
        {
            "figure.facecolor": "#0b1020",
            "axes.facecolor": "#0b1020",
            "savefig.facecolor": "#0b1020",
            "axes.edgecolor": "#475569",
            "axes.labelcolor": "#e5e7eb",
            "xtick.color": "#e5e7eb",
            "ytick.color": "#e5e7eb",
            "text.color": "#f8fafc",
            "axes.titlecolor": "#f8fafc",
            "font.size": 10,
            "axes.grid": True,
            "grid.color": "#334155",
            "grid.alpha": 0.32,
        }
    )


def save_plot(fig, filename: str) -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    path = ASSET_DIR / filename
    fig.savefig(path, dpi=210, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)
    return path


def plot_overview(summary: dict[str, Any]) -> Path:
    setup_plot_style()
    cpu_rows = summary["cpu_rows"]
    gpu_rows = summary["gpu_rows"]
    configs = [
        ("Framsticks EA", cpu_rows, "#f59e0b"),
        ("TreeVAE\nEA ukryta", filter_rows(gpu_rows, label="09_tree_vae_epoch80", method="latent_ea"), "#2dd4bf"),
        ("Transformer\nEA ukryta", filter_rows(gpu_rows, label="10_transformer_vae_epoch90", method="latent_ea"), "#38bdf8"),
        ("TreeVAE\nCMA-ES", filter_rows(gpu_rows, label="09_tree_vae_epoch80", method="cmaes"), "#a78bfa"),
        ("Fitness-aware\nCMA-ES", filter_rows(gpu_rows, label=FITNESS_AWARE, method="cmaes"), "#34d399"),
        ("Oracle portfela\nmetod*", list(best_by_seed(gpu_rows).values()), "#f472b6"),
    ]
    fig, ax = plt.subplots(figsize=(11.5, 5.2))
    data = [[fl(row["best_score"]) for row in rows] for _, rows, _ in configs]
    bp = ax.boxplot(data, patch_artist=True, showfliers=False, medianprops={"color": "#0b1020", "linewidth": 1.5})
    for patch, (_, _, color) in zip(bp["boxes"], configs):
        patch.set_facecolor(color)
        patch.set_alpha(0.78)
        patch.set_edgecolor("#e5e7eb")
    ax.set_xticklabels([name for name, _, _ in configs], rotation=0)
    ax.set_ylabel("Najlepszy vertpos")
    ax.set_title("Rozkład wyników końcowych, 150 seedów; wybrane konfiguracje z pełnych 15")
    ax.text(0.02, -0.24, "*Oracle: najlepsza konfiguracja wybrana osobno dla każdego seeda; analiza potencjału, nie pojedyncza metoda.", transform=ax.transAxes, fontsize=9, color="#adbac9")
    fig.tight_layout()
    return save_plot(fig, "final_overview_boxplots_pl.png")


def plot_winrate_heatmap(summary: dict[str, Any]) -> Path:
    setup_plot_style()
    gpu_rows = summary["gpu_rows"]
    comp = summary["comp"]
    labels = ["09_tree_vae_epoch80", "10_transformer_vae_epoch90", "02_lhs_latent256_seed42_best", "06_lhs_depth_seed42_best", FITNESS_AWARE]
    methods = ["latent_ea", "cem", "cmaes"]
    win = np.zeros((len(labels), len(methods)))
    diff = np.zeros((len(labels), len(methods)))
    for i, label in enumerate(labels):
        for j, method in enumerate(methods):
            item = comp[(label, method)]
            win[i, j] = item["win_rate"]
            diff[i, j] = item["mean_diff"]
    fig, axes = plt.subplots(1, 2, figsize=(12.5, 5.3), sharey=True)
    im0 = axes[0].imshow(win, vmin=0.0, vmax=1.0, cmap="viridis")
    im1 = axes[1].imshow(diff, vmin=-0.6, vmax=0.15, cmap="coolwarm")
    for ax, matrix, title, percent in [
        (axes[0], win, "Win-rate vs Framsticks EA", True),
        (axes[1], diff, "Średnia różnica wyniku", False),
    ]:
        ax.set_xticks(range(len(methods)))
        ax.set_xticklabels([METHOD_NAMES[m] for m in methods], rotation=25, ha="right")
        ax.set_yticks(range(len(labels)))
        ax.set_yticklabels([label_name(label) for label in labels])
        ax.set_title(title)
        ax.grid(False)
        for i in range(len(labels)):
            for j in range(len(methods)):
                value = matrix[i, j]
                text = f"{value * 100:.0f}%" if percent else f"{value:+.3f}"
                ax.text(j, i, text, ha="center", va="center", color="#f8fafc", fontsize=9)
    fig.colorbar(im0, ax=axes[0], shrink=0.8)
    fig.colorbar(im1, ax=axes[1], shrink=0.8)
    fig.tight_layout()
    return save_plot(fig, "final_winrate_heatmap_pl.png")


def plot_bucket_effect(summary: dict[str, Any]) -> Path:
    setup_plot_style()
    cpu_rows = summary["cpu_rows"]
    gpu_rows = summary["gpu_rows"]
    baseline = summary["baseline"]
    oracle_all = best_by_seed(gpu_rows)
    oracle_ea = best_by_seed(gpu_rows, method="latent_ea")

    def bucket_mean(rows: list[dict[str, str]]) -> list[float]:
        grouped: dict[str, list[float]] = defaultdict(list)
        for row in rows:
            grouped[row["bucket_id"]].append(fl(row["best_score"]))
        return [mean(grouped[bucket]) for bucket in BUCKET_ORDER]

    def oracle_diff(oracle: dict[tuple[str, str, str], dict[str, str]]) -> list[float]:
        grouped: dict[str, list[float]] = defaultdict(list)
        for key, row in oracle.items():
            grouped[key[0]].append(fl(row["best_score"]) - fl(baseline[key]["best_score"]))
        return [mean(grouped[bucket]) for bucket in BUCKET_ORDER]

    fig, axes = plt.subplots(1, 2, figsize=(12.5, 4.8))
    x = np.arange(len(BUCKET_ORDER))
    axes[0].plot(x, bucket_mean(cpu_rows), marker="o", color="#f59e0b", linewidth=2.2, label="Framsticks EA")
    axes[0].plot(x, bucket_mean(filter_rows(gpu_rows, label="09_tree_vae_epoch80", method="latent_ea")), marker="o", color="#2dd4bf", linewidth=2.2, label="TreeVAE EA ukryta")
    axes[0].plot(x, bucket_mean(filter_rows(gpu_rows, label="10_transformer_vae_epoch90", method="latent_ea")), marker="o", color="#38bdf8", linewidth=2.2, label="Transformer EA ukryta")
    axes[0].plot(x, bucket_mean(list(oracle_all.values())), marker="o", color="#f472b6", linewidth=2.2, label="Oracle portfela metod")
    axes[0].set_title("Średni wynik końcowy")
    axes[0].set_ylabel("Best vertpos")
    axes[0].legend(fontsize=8)
    axes[1].bar(x - 0.18, oracle_diff(oracle_ea), width=0.36, color="#2dd4bf", label="najlepsza EA ukryta")
    axes[1].bar(x + 0.18, oracle_diff(oracle_all), width=0.36, color="#f472b6", label="oracle portfela metod")
    axes[1].axhline(0.0, color="#e5e7eb", linewidth=1.0)
    axes[1].set_title("Przewaga względem baseline")
    axes[1].set_ylabel("Średnia różnica")
    axes[1].legend(fontsize=8)
    for ax in axes:
        ax.set_xticks(x)
        ax.set_xticklabels([BUCKET_LABELS[bucket] for bucket in BUCKET_ORDER], rotation=25, ha="right")
        ax.set_xlabel("Fitness seeda startowego")
    fig.tight_layout()
    return save_plot(fig, "final_bucket_effect_pl.png")


def x_value(row: dict[str, str]) -> int:
    raw = fl(row.get("evaluations_requested"))
    if not math.isnan(raw) and raw > 0:
        return int(raw)
    return max(1, int(fl(row.get("step") or 0))) * 100


def aggregate_curve(rows: list[dict[str, str]], value_col: str = "best_score") -> tuple[list[int], list[float], list[float], list[float]]:
    grouped: dict[int, list[float]] = defaultdict(list)
    for row in rows:
        grouped[x_value(row)].append(fl(row[value_col]))
    xs = sorted(grouped)
    return (
        xs,
        [median(grouped[x]) for x in xs],
        [quantile(grouped[x], 0.25) for x in xs],
        [quantile(grouped[x], 0.75) for x in xs],
    )


def history_index(rows: list[dict[str, str]]) -> dict[tuple[str, str, str, str, str, str], list[dict[str, str]]]:
    indexed: dict[tuple[str, str, str, str, str, str], list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        indexed[summary_key(row)].append(row)
    for group in indexed.values():
        group.sort(key=x_value)
    return indexed


def histories_for_summaries(summary_rows: dict[tuple[str, str, str], dict[str, str]], indexed: dict[tuple[str, str, str, str, str, str], list[dict[str, str]]]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for row in summary_rows.values():
        rows.extend(indexed.get(summary_key(row), []))
    return rows


def plot_history(summary: dict[str, Any]) -> Path:
    setup_plot_style()
    cpu_hist = summary["cpu_hist"]
    gpu_hist = summary["gpu_hist"]
    gpu_rows = summary["gpu_rows"]
    indexed = history_index(gpu_hist)
    series = [
        ("Framsticks EA", cpu_hist, "#f59e0b"),
        ("TreeVAE EA ukryta", filter_rows(gpu_hist, label="09_tree_vae_epoch80", method="latent_ea"), "#2dd4bf"),
        ("Transformer EA ukryta", filter_rows(gpu_hist, label="10_transformer_vae_epoch90", method="latent_ea"), "#38bdf8"),
        ("TreeVAE CMA-ES", filter_rows(gpu_hist, label="09_tree_vae_epoch80", method="cmaes"), "#a78bfa"),
        ("Oracle portfela metod", histories_for_summaries(best_by_seed(gpu_rows), indexed), "#f472b6"),
    ]
    fig, axes = plt.subplots(1, 2, figsize=(12.7, 4.9))
    for name, rows, color in series:
        xs, med, q25, q75 = aggregate_curve(rows, "best_score")
        axes[0].plot(xs, med, color=color, linewidth=2.0, label=name)
        axes[0].fill_between(xs, q25, q75, color=color, alpha=0.11, linewidth=0)
        xs, med, q25, q75 = aggregate_curve(rows, "best_minus_seed_source")
        axes[1].plot(xs, med, color=color, linewidth=2.0, label=name)
        axes[1].fill_between(xs, q25, q75, color=color, alpha=0.11, linewidth=0)
    axes[0].set_title("Medianowy best-so-far")
    axes[0].set_ylabel("Best vertpos")
    axes[1].set_title("Medianowa poprawa względem seeda")
    axes[1].set_ylabel("Delta")
    axes[1].axhline(0.0, color="#e5e7eb", linewidth=1.0)
    for ax in axes:
        ax.set_xlabel("Ewaluacje true fitness na seed")
    axes[0].legend(fontsize=7, loc="lower right")
    fig.tight_layout()
    return save_plot(fig, "final_history_curves_pl.png")


def plot_budget_100_300(summary: dict[str, Any]) -> Path:
    setup_plot_style()
    full = summary["gpu_label_method_stats"]
    h100 = summary["gpu_label_method_100"]
    baseline_full = summary["cpu_stats"]
    baseline_100 = summary["cpu100_stats"]
    configs = [
        ("Framsticks EA", None, None, "#f59e0b"),
        ("TreeVAE\nEA ukryta", "09_tree_vae_epoch80", "latent_ea", "#2dd4bf"),
        ("Transformer\nEA ukryta", "10_transformer_vae_epoch90", "latent_ea", "#38bdf8"),
        ("TreeVAE\nCMA-ES", "09_tree_vae_epoch80", "cmaes", "#a78bfa"),
    ]
    x = np.arange(len(configs))
    y100 = []
    y300 = []
    for _, label, method, _ in configs:
        if label is None:
            y100.append(baseline_100["mean"])
            y300.append(baseline_full["mean"])
        else:
            y100.append(h100[(label, method)]["mean"])
            y300.append(full[(label, method)]["mean"])
    fig, ax = plt.subplots(figsize=(10.5, 4.8))
    ax.bar(x - 0.18, y100, 0.36, color="#64748b", label="100 generacji")
    ax.bar(x + 0.18, y300, 0.36, color="#2dd4bf", label="300 generacji")
    ax.set_xticks(x)
    ax.set_xticklabels([name for name, _, _, _ in configs])
    ax.set_ylabel("Średni best vertpos")
    ax.set_title("Wpływ budżetu: 100 vs 300 generacji/iteracji")
    for idx, (a, b) in enumerate(zip(y100, y300)):
        ax.text(idx + 0.18, b + 0.025, f"+{b - a:.3f}", ha="center", fontsize=9, color="#f8fafc")
    ax.legend()
    fig.tight_layout()
    return save_plot(fig, "final_budget_100_300_pl.png")


def create_charts(summary: dict[str, Any]) -> dict[str, Path]:
    return {
        "overview": plot_overview(summary),
        "heatmap": plot_winrate_heatmap(summary),
        "buckets": plot_bucket_effect(summary),
        "history": plot_history(summary),
        "budget": plot_budget_100_300(summary),
    }


def add_bg(slide, accent: RGBColor = CYAN) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG
    top = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.035))
    top.fill.solid()
    top.fill.fore_color.rgb = accent
    top.line.fill.background()


def add_text(slide, text: str, x: float, y: float, w: float, h: float, size: int = 18, color: RGBColor = TEXT, bold: bool = False, align: PP_ALIGN | None = None) -> Any:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.margin_left = 0
    p.margin_right = 0
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if align is not None:
        p.alignment = align
    return box


def add_title(slide, title: str, kicker: str, accent: RGBColor = CYAN) -> None:
    add_text(slide, kicker.upper(), 0.65, 0.35, 5.5, 0.25, 9, accent, True)
    add_text(slide, title, 0.65, 0.67, 11.7, 0.55, 27, TEXT, True)
    add_rule(slide, 0.65, 1.28, 12.0, accent)


def add_rule(slide, x: float, y: float, w: float, color: RGBColor = LINE, h: float = 0.012) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_metric(slide, x: float, y: float, w: float, value: str, label: str, accent: RGBColor = CYAN) -> None:
    add_text(slide, value, x, y, w, 0.42, 25, accent, True, PP_ALIGN.CENTER)
    add_rule(slide, x + 0.25, y + 0.48, w - 0.5, LINE, 0.01)
    add_text(slide, label, x, y + 0.58, w, 0.45, 9, MUTED, False, PP_ALIGN.CENTER)


def add_bullets(slide, items: list[str], x: float, y: float, w: float, h: float, size: int = 15, color: RGBColor = TEXT) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(7)
        p.margin_left = Inches(0.1)


def add_table(slide, rows: list[list[str]], x: float, y: float, w: float, h: float, widths: list[float] | None = None, font_size: int = 10) -> None:
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(h))
    table = shape.table
    if widths:
        for idx, width in enumerate(widths):
            table.columns[idx].width = Inches(width)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = BG_ALT if r_idx else RGBColor(31, 41, 55)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Aptos"
                paragraph.font.size = Pt(font_size if r_idx else font_size - 1)
                paragraph.font.bold = r_idx == 0
                paragraph.font.color.rgb = TEXT


def add_picture(slide, path: Path, x: float, y: float, w: float, h: float | None = None) -> None:
    if h is None:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    else:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_two_column_text(slide, left_title: str, left_items: list[str], right_title: str, right_items: list[str], accent: RGBColor) -> None:
    add_text(slide, left_title, 0.85, 1.62, 5.5, 0.36, 16, accent, True)
    add_rule(slide, 0.85, 2.05, 5.3, LINE)
    add_bullets(slide, left_items, 0.9, 2.22, 5.4, 3.9, 14)
    add_text(slide, right_title, 6.85, 1.62, 5.5, 0.36, 16, accent, True)
    add_rule(slide, 6.85, 2.05, 5.3, LINE)
    add_bullets(slide, right_items, 6.9, 2.22, 5.4, 3.9, 14)


def make_presentation(summary: dict[str, Any], charts: dict[str, Path]) -> list[SlideText]:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    notes: list[SlideText] = []

    cpu_stats = summary["cpu_stats"]
    tree_ea = summary["gpu_label_method_stats"][("09_tree_vae_epoch80", "latent_ea")]
    trans_ea = summary["gpu_label_method_stats"][("10_transformer_vae_epoch90", "latent_ea")]
    tree_cma = summary["gpu_label_method_stats"][("09_tree_vae_epoch80", "cmaes")]
    oracle = summary["oracle_all"]
    oracle_ea = summary["oracle_latent_ea"]
    tree_comp = summary["comp"][("09_tree_vae_epoch80", "latent_ea")]
    trans_comp = summary["comp"][("10_transformer_vae_epoch90", "latent_ea")]
    b1_oracle = compute_bucket_oracle(summary, "B1", method=None)
    b3_oracle = compute_bucket_oracle(summary, "B3", method=None)
    b6_oracle = compute_bucket_oracle(summary, "B6", method=None)

    # 1
    slide = prs.slides.add_slide(blank)
    add_bg(slide, CYAN)
    add_text(slide, "Wystąpienie nr 2", 0.75, 0.48, 4.8, 0.35, 12, CYAN, True)
    add_text(
        slide,
        "Porównanie optymalizacji kompatybilnych\npodmian i optymalizacji w przestrzeni\nukrytej dla genotypów Framsticks F1",
        0.75,
        1.08,
        11.5,
        1.75,
        30,
        TEXT,
        True,
    )
    add_text(slide, "Wyniki części autoenkoderowej i operatorów w przestrzeni ukrytej", 0.78, 3.12, 10.8, 0.42, 17, MUTED)
    add_text(slide, "Filip Rosiak", 0.82, 6.28, 2.9, 0.32, 13, TEXT, True)
    add_text(slide, "Promotor: dr hab. inż. Maciej Komosiński", 3.35, 6.28, 5.3, 0.32, 13, MUTED)
    add_text(slide, "Seminarium dyplomowe, 2026", 9.6, 6.28, 2.85, 0.32, 13, MUTED)
    notes.append(SlideText("Metryczka", "Witam, nazywam się Filip Rosiak. Tematem pracy jest porównanie optymalizacji kompatybilnych podmian i optymalizacji w przestrzeni ukrytej dla genotypów Framsticks F1. W tym wystąpieniu skupiam się na wynikach części autoenkoderowej: pokazuję, co zostało zaimplementowane, jak wygląda finalny protokół eksperymentalny oraz co wynika z benchmarku."))

    # 2
    slide = prs.slides.add_slide(blank)
    add_bg(slide, CYAN)
    add_title(slide, "Problem i sprawdzana teza", "kontekst", CYAN)
    add_two_column_text(
        slide,
        "Problem",
        [
            "F1 jest tekstowym opisem struktury organizmu, a nie wektorem liczb.",
            "Mutacja napisu może łatwo zniszczyć poprawną lub użyteczną podstrukturę.",
            "Sama poprawność gramatyczna nie mówi, czy organizm będzie miał dobry vertpos.",
        ],
        "Teza eksperymentalna",
        [
            "Encoder mapuje genotyp do wektora z w przestrzeni ukrytej.",
            "Operatory wykonujemy na z, a wynik dekodujemy do F1 i oceniamy Framsticksem.",
            "Porównujemy to ze standardowym Framsticks EA na tych samych seedach.",
        ],
        CYAN,
    )
    notes.append(SlideText("Problem i teza", "Problemem jest to, że genotyp F1 ma strukturę tekstową i drzewiastą. Standardowe operatory Framsticks są naturalnym punktem odniesienia, ale pytanie brzmi, czy część pracy można przenieść do ciągłej przestrzeni ukrytej autoenkodera. Teza nie jest taka, że przestrzeń ukryta zawsze wygra, tylko że może być konkurencyjnym i czasem lepszym sposobem generowania kandydatów."))

    # 3
    slide = prs.slides.add_slide(blank)
    add_bg(slide, BLUE)
    add_title(slide, "Co udało się wykonać", "implementacja", BLUE)
    add_two_column_text(
        slide,
        "Modele i algorytmy",
        [
            "TreeVAE, TransformerVAE, warianty LHS/depth i fitness-aware.",
            "CEM i CMA-ES w przestrzeni ukrytej jako czarnoskrzynkowe baseline’y.",
            "EA w przestrzeni ukrytej: mutacja Gaussowska, interpolacja, ekstrapolacja, kierunek do lepszych osobników i imigranci losowi.",
        ],
        "Eksperymenty",
        [
            "Każdy kandydat jest oceniany przez Framsticks kryterium vertpos.",
            "Te same genotypy startowe są używane w baseline i metodach w przestrzeni ukrytej.",
            "Seedy podzielono na sześć zakresów fitness, żeby ocenić wpływ punktu startowego.",
        ],
        BLUE,
    )
    notes.append(SlideText("Co wykonano", "Zaimplementowałem modele autoenkoderów oraz metody optymalizacji działające na ich przestrzeni ukrytej. Po stronie modeli są TreeVAE, TransformerVAE i warianty LHS oraz fitness-aware. Po stronie algorytmów są CEM, CMA-ES i własny wariant populacyjnej ewolucji w przestrzeni ukrytej. Najważniejsze eksperymentalnie było zachowanie wspólnych seedów, tej samej funkcji celu i historii przebiegów best-so-far."))

    # 4
    slide = prs.slides.add_slide(blank)
    add_bg(slide, VIOLET)
    add_title(slide, "Protokół finalnego benchmarku", "warunki", VIOLET)
    add_metric(slide, 0.75, 1.55, 1.9, "6", "bucketów fitness", VIOLET)
    add_metric(slide, 2.95, 1.55, 1.9, "25", "seedów na bucket", CYAN)
    add_metric(slide, 5.15, 1.55, 1.9, "300", "generacji/iteracji", AMBER)
    add_metric(slide, 7.35, 1.55, 1.9, "100", "osobników", BLUE)
    add_metric(slide, 9.55, 1.55, 1.9, "vertpos", "funkcja celu", GREEN)
    add_table(
        slide,
        [
            ["Element", "Ustawienie"],
            ["Baseline", "Framsticks mutate/crossOver, te same 150 seedów"],
            ["Metody w przestrzeni ukrytej", "EA w przestrzeni ukrytej, CEM, CMA-ES dla 5 modeli"],
            ["Fitness", "prawdziwy Framsticks vertpos"],
            ["Zakres danych", "genotypy Framsticks F1 z datasetu vertpos"],
            ["Porównanie", "sparowane po bucket, seed_rank i dataset_idx"],
        ],
        1.0,
        3.05,
        11.25,
        2.75,
        [2.3, 8.8],
        11,
    )
    notes.append(SlideText("Protokół", "Finalny benchmark ma 150 seedów, po 25 w sześciu zakresach fitness. Baseline jest liczony raz, bo nie zależy od modelu autoenkodera. Metody w przestrzeni ukrytej są liczone dla pięciu modeli i trzech algorytmów. Porównanie jest sparowane: zestawiamy wyniki dla tego samego genotypu startowego, bucketa i indeksu w zbiorze danych."))

    # 5
    slide = prs.slides.add_slide(blank)
    add_bg(slide, AMBER)
    add_title(slide, "Wynik ogólny: metoda nie dominuje, ale jest konkurencyjna", "wyniki", AMBER)
    add_picture(slide, charts["overview"], 0.65, 1.45, 7.95, 4.55)
    add_metric(slide, 8.9, 1.55, 2.0, fmt(cpu_stats["mean"], 3), "średnia baseline", AMBER)
    add_metric(slide, 10.75, 1.55, 2.0, fmt(trans_ea["mean"], 3), "Transformer EA", BLUE)
    add_metric(slide, 8.9, 2.95, 2.0, fmt(tree_ea["mean"], 3), "TreeVAE EA", CYAN)
    add_metric(slide, 10.75, 2.95, 2.0, fmt(tree_cma["max"], 3), "rekord TreeVAE+CMA-ES", VIOLET)
    add_text(slide, "Wykres pokazuje wybrane konfiguracje z pełnych 15: dwie najlepsze EA w przestrzeni ukrytej i dwa warianty CMA-ES. Najlepszy pojedynczy wynik nadal ma Framsticks EA, ale najlepsze konfiguracje EA mają wyższą średnią niż baseline.", 8.9, 4.55, 3.55, 1.45, 12, MUTED)
    notes.append(SlideText("Wynik ogólny", f"W pełnym benchmarku baseline osiąga średnio {fmt(cpu_stats['mean'], 3)}, a najlepszy pojedynczy wynik {fmt(cpu_stats['max'], 3)}. Dwie najważniejsze konfiguracje w przestrzeni ukrytej, TreeVAE i TransformerVAE z operatorami ewolucyjnymi, mają średnie około {fmt(tree_ea['mean'], 3)} i {fmt(trans_ea['mean'], 3)}. To nie jest dominacja, ale jest to realna konkurencyjność wobec standardowych operatorów Framsticks."))

    # 6
    slide = prs.slides.add_slide(blank)
    add_bg(slide, GREEN)
    add_title(slide, "Porównanie sparowane z baseline’em", "wyniki", GREEN)
    add_picture(slide, charts["heatmap"], 0.65, 1.38, 8.3, 4.75)
    add_metric(slide, 9.35, 1.6, 2.15, f"{int(tree_comp['wins'])}/150", "TreeVAE EA wygrywa", CYAN)
    add_metric(slide, 9.35, 3.0, 2.15, f"{int(trans_comp['wins'])}/150", "Transformer EA wygrywa", BLUE)
    add_metric(slide, 9.35, 4.4, 2.15, fmt(oracle_ea["mean_diff"], 3), "oracle EA vs baseline", PINK)
    notes.append(SlideText("Porównanie sparowane", "Ten slajd pokazuje porównanie per seed. Dla każdego seeda wynik danej konfiguracji w przestrzeni ukrytej jest zestawiany z wynikiem baseline’u na tym samym genotypie startowym. Najlepsze konkretne konfiguracje to TreeVAE i TransformerVAE z EA w przestrzeni ukrytej. Oracle nie jest metodą do raportowania jako główny wynik, ale pokazuje potencjał rodziny metod: gdyby umieć dobrać konfigurację do seeda, przewaga byłaby wyraźna."))

    # 7
    slide = prs.slides.add_slide(blank)
    add_bg(slide, CYAN)
    add_title(slide, "Dynamika: przestrzeń ukryta potrzebuje budżetu", "przebiegi", CYAN)
    add_picture(slide, charts["history"], 0.65, 1.38, 8.2, 4.65)
    add_picture(slide, charts["budget"], 8.9, 1.65, 3.45, 2.45)
    add_text(slide, "Po 100 generacjach EA w przestrzeni ukrytej jest blisko baseline’u, ale jeszcze zwykle minimalnie przegrywa. Między 100 a 300 generacją TreeVAE i TransformerVAE nadrabiają więcej niż baseline.", 9.05, 4.35, 3.15, 1.25, 12, MUTED)
    notes.append(SlideText("Dynamika", "Przebiegi pokazują, że baseline Framsticks jest silny wcześnie, ale metody w przestrzeni ukrytej lepiej wykorzystują dłuższy budżet w wybranych konfiguracjach. Po 100 generacjach TreeVAE i TransformerVAE z EA są już blisko baseline’u, a po 300 generacjach ich średnia lekko go przekracza. To sugeruje, że przestrzeń ukryta nie jest tylko szybkim heurystycznym skokiem, ale wymaga czasu na eksplorację."))

    # 8
    slide = prs.slides.add_slide(blank)
    add_bg(slide, PINK)
    add_title(slide, "Największy zysk: słabe i średnie seedy", "analiza bucketów", PINK)
    add_picture(slide, charts["buckets"], 0.65, 1.38, 8.65, 4.75)
    add_metric(slide, 9.7, 1.6, 2.0, fmt(b1_oracle["mean_diff"], 3), "B1 oracle", PINK)
    add_metric(slide, 9.7, 3.0, 2.0, fmt(b3_oracle["mean_diff"], 3), "B3 oracle", CYAN)
    add_metric(slide, 9.7, 4.4, 2.0, fmt(b6_oracle["mean_diff"], 3), "B6 oracle", AMBER)
    notes.append(SlideText("Buckety", "Podział na buckety daje jeden z ciekawszych wniosków. Przestrzeń ukryta najbardziej pomaga dla słabych i średnich punktów startowych, szczególnie w B1 i B3. Dla elit, czyli B6, baseline Framsticks jest trudniejszy do pobicia. To ma sens: gdy start jest już bardzo dobry, lokalne operatory Framsticks potrafią skutecznie dopracować morfologię."))

    # 9
    slide = prs.slides.add_slide(blank)
    add_bg(slide, BLUE)
    add_title(slide, "Co mówią modele i algorytmy", "wnioski z wyników", BLUE)
    add_table(
        slide,
        [
            ["Obserwacja", "Wniosek"],
            ["CEM ma najniższą średnią i win-rate", "prosty rozkład diagonalny za słabo eksploruje użyteczne okolice dekodera"],
            ["CMA-ES daje wysokie maksima", "adaptacja kowariancji pomaga, ale stabilność jest niższa niż w EA"],
            ["TreeVAE i TransformerVAE + EA są najlepsze średnio", "operatory populacyjne w z są najbardziej odporne na seed"],
            ["LHS/depth dobrze rekonstruują, ale nie wygrywają", "rekonstrukcja nie wystarcza jako kryterium wyboru modelu"],
            ["Fitness-aware ma mocne CMA-ES, ale nie najlepszą średnią", "sygnał fitness pomaga, lecz nie rozwiązuje geometrii przestrzeni"],
        ],
        0.85,
        1.55,
        11.7,
        4.55,
        [4.2, 7.5],
        11,
    )
    notes.append(SlideText("Modele i algorytmy", "Wyniki nie mówią tylko, która liczba jest największa. Pokazują też, które mechanizmy są stabilne. CEM jest słaby jako główna metoda. CMA-ES potrafi dawać bardzo dobre maksima, ale jest mniej równy. Najbardziej stabilny wariant to EA w przestrzeni ukrytej dla TreeVAE i TransformerVAE. Ważny negatywny wniosek jest taki, że dobre rekonstrukcje modeli LHS nie przekładają się automatycznie na dobrą optymalizację."))

    # 10
    slide = prs.slides.add_slide(blank)
    add_bg(slide, RED)
    add_title(slide, "Ograniczenia i wiarygodność porównania", "uczciwa ocena", RED)
    add_two_column_text(
        slide,
        "Ograniczenia",
        [
            "Nie ma jednej konfiguracji, która wygrywa zawsze z baseline’em.",
            "Oracle pokazuje potencjał portfela metod, a nie algorytm możliwy do użycia bez wyboru konfiguracji.",
            "Wyniki zależą od zakresu fitness seeda, więc sama średnia nie opisuje całego zachowania.",
        ],
        "Co zrobiono dodatkowo",
        [
            "Zastosowano sparowane porównanie: ten sam seed startowy dla baseline’u i metod ukrytych.",
            "Dodano pełne historie best-so-far, co pozwala porównywać metody po budżecie ewaluacji.",
            "Wyniki są raportowane dla 150 seedów, nie tylko dla pojedynczych rekordów.",
        ],
        RED,
    )
    notes.append(SlideText("Ograniczenia", "Najważniejsze ograniczenie interpretacyjne jest takie, że nie ma jednej konfiguracji wygrywającej zawsze. Dlatego wyniki trzeba czytać na trzech poziomach: konkretna metoda, najlepsza metoda w rodzinie oraz oracle portfela metod. Druga rzecz to stabilność: pojedynczy rekord nie wystarcza, dlatego prezentuję wyniki po 150 seedach i z podziałem na zakresy fitness punktu startowego."))

    # 11
    slide = prs.slides.add_slide(blank)
    add_bg(slide, GREEN)
    add_title(slide, "Wnioski", "podsumowanie", GREEN)
    add_metric(slide, 0.9, 1.55, 2.0, fmt(tree_ea["mean"] - cpu_stats["mean"], 3), "TreeVAE EA vs baseline", CYAN)
    add_metric(slide, 3.15, 1.55, 2.0, fmt(trans_ea["mean"] - cpu_stats["mean"], 3), "Transformer EA vs baseline", BLUE)
    add_metric(slide, 5.4, 1.55, 2.0, f"{int(oracle['wins'])}/150", "oracle wygrywa", PINK)
    add_metric(slide, 7.65, 1.55, 2.0, fmt(oracle["mean_diff"], 3), "oracle średnio vs baseline", PINK)
    add_metric(slide, 9.9, 1.55, 2.0, fmt(cpu_stats["max"], 3), "najlepszy rekord baseline", AMBER)
    add_bullets(
        slide,
        [
            "EA w przestrzeni ukrytej jest realnie konkurencyjna ze standardowym Framsticks EA, ale przewaga zależy od modelu i seeda.",
            "Najbardziej użyteczny efekt widać dla słabych i średnich seedów; dla elit baseline pozostaje bardzo mocny.",
            "CMA-ES jest dobry do szukania rekordów, ale EA w przestrzeni ukrytej jest stabilniejsza średnio.",
            "Wniosek negatywny: sama rekonstrukcja i sama gramatyczność generowania nie wystarczają do wysokiego fitness.",
        ],
        1.0,
        3.2,
        11.55,
        2.7,
        15,
    )
    notes.append(SlideText("Wnioski", "Główny wniosek jest pozytywny, ale nie bezwarunkowy. Udało się pokazać, że operatory w przestrzeni ukrytej mogą być konkurencyjne wobec standardowych operatorów Framsticks. Najlepsze konkretne konfiguracje mają średnią lekko wyższą od baseline’u, ale najlepszy pojedynczy rekord nadal należy do baseline’u. Najważniejsze jest to, że efekt jest zależny od zakresu fitness seeda: przestrzeń ukryta najbardziej pomaga tam, gdzie start jest słaby albo średni."))

    # 12
    slide = prs.slides.add_slide(blank)
    add_bg(slide, VIOLET)
    add_title(slide, "Plany rozwojowe", "następne kroki", VIOLET)
    add_bullets(
        slide,
        [
            "Re-ewaluacja najlepszych genotypów znalezionych w przestrzeni ukrytej i oznaczenie novelty względem datasetu.",
            "Adaptacyjne operatory w przestrzeni ukrytej: skala mutacji i imigranci zależni od plateau przebiegu.",
            "Dobór konfiguracji do bucketa fitness, bo wyniki pokazują różne zachowanie dla słabych i elitarnych seedów.",
            "Retraining lub fine-tuning autoenkodera na elitach znalezionych przez przeszukiwanie.",
            "Końcowe porównanie z CoSO w jednolitym protokole raportowania.",
        ],
        0.95,
        1.55,
        11.7,
        4.65,
        17,
    )
    notes.append(SlideText("Plany", "Najbliższe kroki wynikają bezpośrednio z wyników. Po pierwsze, trzeba reewaluować najlepsze genotypy znalezione przez metody w przestrzeni ukrytej. Po drugie, warto dopasować operatory do przebiegu, bo wykresy pokazują plateau. Po trzecie, osobnym kierunkiem jest dobór metody do bucketa fitness, bo dla różnych punktów startowych najlepsze są inne strategie."))

    # 13
    slide = prs.slides.add_slide(blank)
    add_bg(slide, AMBER)
    add_title(slide, "Referencje", "źródła", AMBER)
    refs = [
        "Kingma, Welling: Auto-Encoding Variational Bayes, ICLR 2014.",
        "Kusner, Paige, Hernández-Lobato: Grammar Variational Autoencoder, 2017.",
        "Kaszuba, Komosiński, Mensfelt: optymalizacja sekwencji z użyciem reprezentacji autoenkoderowych, IEEE CEC 2021.",
        "Hansen: The CMA Evolution Strategy, tutorial.",
        "Framsticks documentation and FramsticksLib API; kryterium: vertpos.",
        "Dataset F1: genotypy Framsticks F1 oceniane kryterium vertpos.",
    ]
    add_bullets(slide, refs, 0.95, 1.55, 11.7, 4.45, 15)
    add_text(slide, "Dziękuję za uwagę", 0.95, 6.35, 4.2, 0.4, 22, CYAN, True)
    notes.append(SlideText("Referencje", "Na końcu podaję podstawowe źródła: VAE, Grammar VAE, prace o optymalizacji sekwencji w przestrzeni ukrytej, CMA-ES oraz dokumentację Framsticks. Dziękuję za uwagę."))

    prs.save(PPTX_PATH)
    return notes


def compute_bucket_oracle(summary: dict[str, Any], bucket: str, method: str | None) -> dict[str, float]:
    gpu_rows = [row for row in summary["gpu_rows"] if row["bucket_id"] == bucket]
    baseline = {key: row for key, row in summary["baseline"].items() if key[0] == bucket}
    oracle = best_by_seed(gpu_rows, method=method)
    diffs: list[float] = []
    scores: list[float] = []
    for key, row in oracle.items():
        base = baseline[key]
        scores.append(fl(row["best_score"]))
        diffs.append(fl(row["best_score"]) - fl(base["best_score"]))
    return {
        "n": float(len(scores)),
        "wins": float(sum(diff > 1e-9 for diff in diffs)),
        "mean": mean(scores),
        "median": median(scores),
        "max": max(scores),
        "mean_diff": mean(diffs),
        "median_diff": median(diffs),
    }


def write_notes(notes: list[SlideText], summary: dict[str, Any]) -> None:
    lines = ["# Wystąpienie nr 2 - tekst do slajdów", ""]
    for idx, slide in enumerate(notes, start=1):
        lines.append(f"## Slajd {idx}. {slide.title}")
        lines.append("")
        lines.append(slide.speaker)
        lines.append("")

    lines.extend(
        [
            "## Dodatkowe wnioski do wykorzystania w dyskusji",
            "",
            "To nie jest część do czytania slajd po slajdzie. To zestaw wniosków i interpretacji, które można wykorzystać przy pytaniach albo przy pisaniu rozdziału z wynikami.",
            "",
            "1. Przestrzeń ukryta nie zastępuje automatycznie operatorów Framsticks, ale tworzy inną klasę ruchów w przestrzeni rozwiązań. Najlepiej widać to dla słabszych i średnich seedów, gdzie klasyczne lokalne mutacje nie zawsze szybko trafiają w dobrą makrostrukturę.",
            "2. Wyniki sugerują efekt zależny od punktu startowego. Dla bucketów B1 i B3 oracle portfela metod wygrywa wyraźnie, natomiast dla B6 baseline lepiej wykorzystuje już dobre genotypy. To uzasadnia przyszły dobór strategii do zakresu fitness seeda.",
            "3. CEM jest ważnym wynikiem negatywnym. Sam fakt, że optymalizujemy w przestrzeni ukrytej, nie wystarcza. Zbyt prosta aktualizacja rozkładu nie wykorzystuje dobrze geometrii dekodera.",
            "4. CMA-ES jest dobry do rekordów, ale mniej stabilny niż EA w przestrzeni ukrytej. W pracy można go traktować jako mocny czarnoskrzynkowy baseline dla przestrzeni ukrytej, a nie jako bezpośredni zamiennik operatorów ewolucyjnych.",
            "5. Najbardziej obiecująca pojedyncza metoda to EA w przestrzeni ukrytej dla TreeVAE lub TransformerVAE. Obie konfiguracje mają średnią lekko powyżej baseline’u i dodatni wynik w porównaniu sparowanym, ale nie wygrywają każdego seeda.",
            "6. Najlepszy rekonstruktor nie jest automatycznie najlepszym optymalizatorem. To ważny argument metodologiczny: model do przeszukiwania trzeba wybierać po właściwościach generatywnych przestrzeni ukrytej, nie tylko po rekonstrukcji walidacyjnej.",
            "7. Wykres 100 vs 300 generacji pokazuje, że część przewagi metod w przestrzeni ukrytej pojawia się dopiero przy dłuższym budżecie. To sugeruje, że problemem nie jest tylko jakość pojedynczej mutacji, ale dynamika populacji i dywersyfikacja.",
            "8. Sparowanie seedów było konieczne metodologicznie. Bez tego różnice między metodami mieszałyby efekt algorytmu z trudnością punktu startowego.",
            "9. Oracle portfela metod nie jest metodą, ale pokazuje potencjalną wartość portfela strategii. Jeśli w przyszłości uda się przewidywać, która konfiguracja pasuje do seeda, można zbliżyć się do wyników oracle bez ręcznego wyboru po fakcie.",
            "10. Do finalnej pracy warto oddzielić trzy poziomy twierdzeń: konkretna metoda, najlepsza metoda w rodzinie dla danego modelu, oraz oracle jako analiza potencjału. Mieszanie tych poziomów prowadziłoby do zbyt optymistycznych wniosków.",
            "",
            "## Liczby kontrolne z aktualnego benchmarku",
            "",
            f"- Framsticks EA baseline: {len(summary['cpu_rows'])} runów, mean={fmt(summary['cpu_stats']['mean'])}, median={fmt(summary['cpu_stats']['median'])}, max={fmt(summary['cpu_stats']['max'])}.",
            f"- Metody w przestrzeni ukrytej: {len(summary['gpu_rows'])} runów, 5 modeli × 3 metody × 6 bucketów × 25 seedów.",
            f"- TreeVAE + EA w przestrzeni ukrytej: mean={fmt(summary['gpu_label_method_stats'][('09_tree_vae_epoch80', 'latent_ea')]['mean'])}, win-rate={pct(summary['comp'][('09_tree_vae_epoch80', 'latent_ea')]['win_rate'])}.",
            f"- TransformerVAE + EA w przestrzeni ukrytej: mean={fmt(summary['gpu_label_method_stats'][('10_transformer_vae_epoch90', 'latent_ea')]['mean'])}, win-rate={pct(summary['comp'][('10_transformer_vae_epoch90', 'latent_ea')]['win_rate'])}.",
            f"- Oracle portfela metod: wins={int(summary['oracle_all']['wins'])}/150, mean_diff={fmt(summary['oracle_all']['mean_diff'])}, max={fmt(summary['oracle_all']['max'])}.",
            "",
        ]
    )
    NOTES_PATH.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    summary = compute_summary()
    charts = create_charts(summary)
    notes = make_presentation(summary, charts)
    write_notes(notes, summary)
    print(f"Wrote: {PPTX_PATH}")
    print(f"Wrote: {NOTES_PATH}")
    print(f"Wrote assets: {ASSET_DIR}")


if __name__ == "__main__":
    main()
